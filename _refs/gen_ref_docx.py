# -*- coding: utf-8 -*-
# 生成毕业设计参考文献 Word 文件（20 篇，GB/T 7714—2015 格式）
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

REFS = [
    # 中文期刊论文 [J]
    "丁禹钧, 朱一龙, 王雪静. 基于SpringBoot和Vue的高校学生社团管理系统[J]. 电脑编程技巧与维护, 2025(9): 110-112, 165.",
    "陈玲. 基于SpringBoot和Vue3的高校社团管理系统设计与实现[J]. 电脑编程技巧与维护, 2025(8): 3-7, 23.",
    "李建超, 陈平. 基于微服务架构的高校多角色社团协同管理平台的设计与实现[J]. 电脑知识与技术, 2026(13): 41-44.",
    "孙梦菲, 周天源, 王天澍. 基于微信小程序的校园社团管理系统设计与实现[J]. 现代信息科技, 2024(3): 40-44.",
    "苏永辉. 基于协同过滤算法的校园社团推荐软件的设计与应用[J]. 黑龙江科学, 2024(13): 76-79.",
    "张浩. 人工智能驱动下高校创新创业孵化平台设计与实现[J]. 数字技术与应用, 2025(7): 51-53.",
    "赵芳芳. 应用型高校构建全链条创新创业孵化模式研究[J]. 黄河科技学院学报, 2025(7): 88-95.",
    "孔晓茵, 张建双, 张娇娇. 立德树人视域下高校学生社团管理的创新路径[J]. 西部学刊, 2026(3): 92-95.",
    "梁宵, 梁本部. 高校学生社团管理风险的影响因素与作用路径——基于扎根理论与DEMATEL-ISM的分析[J]. 武汉理工大学学报(社会科学版), 2026(4): 38-47.",
    # 中文学位论文 [D]
    "燕达. 基于VUE的可视化智慧校园系统的设计与实现[D]. 天津: 天津师范大学, 2023.",
    "牛子逸. 基于Vue+SpringBoot的音乐评阅系统设计与实现[D]. 成都: 电子科技大学, 2025.",
    "唐双林. 基于Vue和SpringBoot架构的智能推荐农产品团购销售系统[D]. 重庆: 重庆三峡学院, 2023.",
    "朱珊珮. 中职学校学生社团管理现状、问题及对策研究[D]. 武汉: 湖北大学, 2024.",
    # 中文专著/教材 [M]
    "李磊. Java EE企业级应用开发实战（Spring Boot+Vue+Element）[M]. 北京: 人民邮电出版社, 2023.",
    "十三, 尼克陈. Spring Boot+Vue 3大型前后端分离项目实战[M]. 北京: 电子工业出版社, 2023.",
    "王红元, 刘军. Vue.js 3+TypeScript完全指南[M]. 北京: 电子工业出版社, 2023.",
    "杨海民. Vue.js 3高级编程：UI组件库开发实战[M]. 北京: 电子工业出版社, 2024.",
    # 英文期刊论文 [J]
    "QI W L. Design and implementation of university club management system based on WeChat Applet[J]. Journal of Computing and Electronic Information Management, 2024, 13(1): 25-29.",
    "LASUT A H, SAHULATA R A, SILITONGA A O, et al. Web-based information system for student organization management at Universitas Klabat[J]. CogITo Smart Journal, 2026, 12(1): 160-171.",
    "HUANG H J, LI B Y. Design and implementation of student management system of integrated programmable device programming system[J]. Scientific Reports, 2024, 14(1): 11873.",
]


def set_font(run, size=10.5, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")


doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(3.0)

h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(h.add_run("参考文献"), size=16, bold=True)

for i, ref in enumerate(REFS, 1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.85)
    pf.first_line_indent = Cm(-0.85)   # 悬挂缩进
    pf.space_after = Pt(3)
    pf.line_spacing = 1.25
    set_font(p.add_run("[%d] %s" % (i, ref)))

out = r"C:\Users\86187\Desktop\参考文献（20篇·GB-T 7714）.docx"
doc.save(out)
print("SAVED:", out)
